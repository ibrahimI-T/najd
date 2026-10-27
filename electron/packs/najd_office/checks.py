"""Checks and layout-aware reading without Office: lint(path) for decks and inspect(path) for decks, documents and
workbooks. (The external-content scan that guards the Office route lives in NAJD itself, electron/office/office-scan.cjs.)"""
import json
import re

from . import motion
from .text import font_installed, is_arabic, measure_lines, measure_width

EMU = 914400
SHEET_XML_BYTES = 64 * 1024 * 1024
MAX_CELLS = 200000
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _rgb(value):
    value = str(value or "")
    return tuple(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4)) if re.fullmatch(r"[0-9A-Fa-f]{6}", value) else None


def contrast(a, b):
    def luminance(rgb):
        channel = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
        r, g, b = (channel(c) for c in rgb)
        return 0.2126 * r + 0.7152 * g + 0.0722 * b
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _solid(fill_owner):
    try:
        if fill_owner.fill.type == 1:  # MSO_FILL.SOLID
            return str(fill_owner.fill.fore_color.rgb)
    except Exception:
        return None
    return None


def _box(shape):
    return (shape.left or 0, shape.top or 0, (shape.left or 0) + (shape.width or 0), (shape.top or 0) + (shape.height or 0))


def _inside(inner, outer, slack=EMU * 0.02):
    return inner[0] >= outer[0] - slack and inner[1] >= outer[1] - slack and inner[2] <= outer[2] + slack and inner[3] <= outer[3] + slack


def _overlap(a, b):
    width = min(a[2], b[2]) - max(a[0], b[0])
    height = min(a[3], b[3]) - max(a[1], b[1])
    return width * height if width > 0 and height > 0 else 0


def _runs(shape):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            yield paragraph, run


def lint(path):
    """Issues a reviewer would catch on the slides, found without Office. Text fit is estimated from the real font's
    metrics; animations are reported as declared in the file, not as played."""
    lower = str(path).lower()
    if not lower.endswith(".pptx"):
        # Slides only: a document or workbook gets a note instead of a failure, so a task does not spend a step on the error
        # (owner live test 2026-10-25). inspect() and najd_office_review check documents and workbooks.
        return {"kind": lower.rsplit(".", 1)[-1] if "." in lower else "unknown", "issues": [], "slides": [],
                "note": "lint checks slides (.pptx) only; for .docx and .xlsx use inspect(path) or najd_office_review"}
    from pptx import Presentation
    prs = Presentation(path)
    width, height = prs.slide_width, prs.slide_height
    issues, slides = [], []
    for number, slide in enumerate(prs.slides, 1):
        background = None
        try:
            if slide.background.fill.type == 1:
                background = str(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        shapes = list(slide.shapes)
        ids = {shape.shape_id for shape in shapes}
        texts = []
        for shape in shapes:
            box = _box(shape)
            if box[0] < -EMU * 0.01 or box[1] < -EMU * 0.01 or box[2] > width + EMU * 0.01 or box[3] > height + EMU * 0.01:
                issues.append({"slide": number, "shape": shape.name, "issue": "off-slide"})
            if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
                issues.append({"slide": number, "shape": shape.name, "issue": "empty placeholder"})
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            texts.append((shape, box))
            frame = shape.text_frame
            inner_w = ((shape.width or 0) - (frame.margin_left or 0) - (frame.margin_right or 0)) / 12700
            inner_h = ((shape.height or 0) - (frame.margin_top or 0) - (frame.margin_bottom or 0)) / 12700
            needed = 0
            fill = _solid(shape) or background
            for paragraph in frame.paragraphs:
                runs = paragraph.runs
                if not runs:
                    continue
                size = (runs[0].font.size.pt if runs[0].font.size else 18)
                family = runs[0].font.name or "Calibri"
                spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else 1.0
                text = "".join(run.text for run in runs)
                if frame.word_wrap is False:
                    # Without wrapping a line is as wide as its text (M3-I026).
                    lines = len(text.split("\n"))
                    if measure_width(text, family, size, bool(runs[0].font.bold)) > inner_w * 1.02:
                        issues.append({"slide": number, "shape": shape.name, "issue": "text wider than its box (wrapping is off)"})
                else:
                    lines = measure_lines(text, family, size, max(inner_w, 1), bool(runs[0].font.bold))
                needed += lines * size * 1.2 * spacing
                for run in runs:
                    if run.font.name and not font_installed(run.font.name):
                        issues.append({"slide": number, "shape": shape.name, "issue": "font not installed: " + run.font.name})
                    try:
                        color = str(run.font.color.rgb) if run.font.color and run.font.color.type is not None else None
                    except Exception:
                        color = None
                    if color and fill and _rgb(color) and _rgb(fill):
                        ratio = contrast(_rgb(color), _rgb(fill))
                        large = size >= 18 or (size >= 14 and run.font.bold)
                        if ratio < (3.0 if large else 4.5):
                            issues.append({"slide": number, "shape": shape.name, "issue": "low contrast %.1f:1" % ratio})
                    if is_arabic(run.text) and paragraph._p.pPr is not None and paragraph._p.pPr.get("rtl") != "1":
                        issues.append({"slide": number, "shape": shape.name, "issue": "Arabic paragraph not right-to-left"})
            # Auto-fit only resizes when PowerPoint edits the text, not when it shows the slide, so it does not excuse overflow.
            if needed > inner_h * 1.05:
                issues.append({"slide": number, "shape": shape.name, "issue": "text overflows (needs %.0f pt of %.0f pt)" % (needed, inner_h)})
        # Text boxes that partly overlap each other (a text inside a card is fine; a card under text is fine).
        for index, (a, box_a) in enumerate(texts):
            for b, box_b in texts[index + 1:]:
                area = _overlap(box_a, box_b)
                smaller = min((box_a[2] - box_a[0]) * (box_a[3] - box_a[1]), (box_b[2] - box_b[0]) * (box_b[3] - box_b[1])) or 1
                if area / smaller > 0.1 and not _inside(box_a, box_b) and not _inside(box_b, box_a):
                    issues.append({"slide": number, "shape": a.name + " / " + b.name, "issue": "text overlaps"})
        declared = motion.read(slide._element)
        for effect in declared["effects"]:
            if effect["spid"] not in ids:
                issues.append({"slide": number, "shape": str(effect["spid"]), "issue": "animation targets a missing shape"})
        slides.append({"slide": number, "shapes": len(shapes), "animations": len(declared["effects"]), "transition": (declared["transition"] or {}).get("kind")})
    return {"issues": issues, "slides": slides, "motion": "declared in the file, not played"}


def inspect(path, limit=300):
    """Layout-aware structure of a .pptx, .docx or .xlsx (bounded)."""
    lower = str(path).lower()
    clip = lambda value: (value[:limit] + "…") if isinstance(value, str) and len(value) > limit else value
    if lower.endswith(".pptx"):
        from pptx import Presentation
        prs = Presentation(path)
        out = {"kind": "pptx", "size_in": [round(prs.slide_width / EMU, 2), round(prs.slide_height / EMU, 2)], "slides": []}
        for number, slide in enumerate(prs.slides, 1):
            names = {shape.shape_id: shape.name for shape in slide.shapes}
            shapes = []
            for shape in list(slide.shapes)[:60]:
                row = {"id": shape.shape_id, "name": shape.name, "type": str(shape.shape_type).split(".")[-1].split(" ")[0] if shape.shape_type else "shape",
                       "box_in": [round((value or 0) / EMU, 2) for value in (shape.left, shape.top, shape.width, shape.height)]}
                if shape.has_text_frame and shape.text_frame.text.strip():
                    row["text"] = clip(shape.text_frame.text)
                    fonts = sorted({run.font.name for _, run in _runs(shape) if run.font.name})
                    if fonts:
                        row["fonts"] = fonts
                    if any(p._p.pPr is not None and p._p.pPr.get("rtl") == "1" for p in shape.text_frame.paragraphs):
                        row["rtl"] = True
                if getattr(shape, "has_chart", False) and shape.has_chart:
                    row["chart"] = str(shape.chart.chart_type).split(".")[-1].split(" ")[0]
                if getattr(shape, "has_table", False) and shape.has_table:
                    row["table"] = [len(shape.table.rows), len(shape.table.columns)]
                shapes.append(row)
            declared = motion.read(slide._element)
            for effect in declared["effects"]:
                effect["shape"] = names.get(effect["spid"])
            out["slides"].append({"slide": number, "shapes": shapes, "animations": declared["effects"], "transition": declared["transition"],
                                  "notes": clip(slide.notes_slide.notes_text_frame.text) if slide.has_notes_slide else None})
        return out
    if lower.endswith(".docx"):
        import docx
        document = docx.Document(path)
        paragraphs = []
        for paragraph in document.paragraphs[:200]:
            if not paragraph.text.strip():
                continue
            bidi = paragraph._p.pPr is not None and paragraph._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None
            paragraphs.append({"style": paragraph.style.name if paragraph.style is not None else None, "text": clip(paragraph.text), **({"rtl": True} if bidi else {})})
        section = document.sections[0] if document.sections else None
        return {"kind": "docx", "paragraphs": paragraphs, "tables": [[len(t.rows), len(t.columns)] for t in document.tables],
                "images": len(document.inline_shapes), "page_in": [round(section.page_width / EMU, 2), round(section.page_height / EMU, 2)] if section else None}
    if lower.endswith(".xlsx"):
        return _inspect_xlsx(path, clip)
    raise ValueError("inspect reads .pptx, .docx and .xlsx")


SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _cell_position(ref):
    match = re.match(r"^([A-Z]{1,3})(\d+)$", ref or "")
    if not match:
        return None
    column = 0
    for letter in match.group(1):
        column = column * 26 + ord(letter) - 64
    return int(match.group(2)), column


def _inspect_xlsx(path, clip):
    """Workbook structure read as a bounded stream (M3-I019/I023): nothing builds a workbook in memory, so merged ranges,
    hyperlink ranges and declared dimensions cannot expand. A workbook whose parts (pictures aside) exceed SHEET_XML_BYTES
    uncompressed is skipped; at most MAX_CELLS cells are visited in all; the first cells shown are at most 8 x 26."""
    import posixpath
    import zipfile
    from xml.etree.ElementTree import iterparse

    with zipfile.ZipFile(path) as package:
        total = sum(info.file_size for info in package.infolist() if not info.filename.lower().startswith("xl/media/"))
        if total > SHEET_XML_BYTES:
            return {"kind": "xlsx", "skipped": "the workbook is too large to inspect here (%d MB uncompressed)" % (total // (1024 * 1024))}
        names = set(package.namelist())

        def relationships(part):
            folder, base = posixpath.split(part)
            rels = posixpath.join(folder, "_rels", base + ".rels")
            out = {}
            if rels not in names:
                return out
            with package.open(rels) as stream:
                for _, element in iterparse(stream):
                    if element.tag == PKG_REL + "Relationship" and element.get("TargetMode") != "External":
                        target = element.get("Target") or ""
                        resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
                        out[element.get("Id")] = (element.get("Type") or "", resolved)
            return out

        root = relationships("")
        workbook = next((target for kind, target in root.values() if kind.endswith("/officeDocument")), "xl/workbook.xml")
        if workbook not in names:
            raise ValueError("not a workbook")
        book_rels = relationships(workbook)
        sheets_declared = []
        with package.open(workbook) as stream:
            for _, element in iterparse(stream):
                if element.tag == SHEET_NS + "sheet":
                    sheets_declared.append((element.get("name"), element.get(REL_NS + "id")))
        strings_part = next((target for kind, target in book_rels.values() if kind.endswith("/sharedStrings")), None)

        visited, sheets, wanted = 0, [], {}
        for name, rid in sheets_declared[:20]:
            part = book_rels.get(rid, ("", None))[1]
            if not part or part not in names:
                continue
            formulas, grid, dimension, complete, max_row, max_col = [], {}, None, True, 0, 0
            with package.open(part) as stream:
                for _, element in iterparse(stream):
                    tag = element.tag
                    if tag == SHEET_NS + "dimension":
                        dimension = element.get("ref")
                    elif tag == SHEET_NS + "c":
                        visited += 1
                        if visited > MAX_CELLS:
                            complete = False
                            break
                        ref = element.get("r")
                        formula = element.find(SHEET_NS + "f")
                        if formula is not None:
                            formulas.append(ref)
                        position = _cell_position(ref)
                        if position and position[0] <= 8 and position[1] <= 26:
                            kind = element.get("t")
                            value = element.find(SHEET_NS + "v")
                            if formula is not None:
                                text = "=" + (formula.text or "")
                            elif kind == "s" and value is not None and (value.text or "").isdigit():
                                text = ("s", int(value.text))
                                wanted[int(value.text)] = None
                            elif kind == "inlineStr":
                                text = "".join(node.text or "" for node in element.iter(SHEET_NS + "t"))
                            elif kind == "b" and value is not None:
                                text = "TRUE" if value.text == "1" else "FALSE"
                            else:
                                text = value.text if value is not None else None
                            if text is not None:
                                grid[position] = text
                                max_row, max_col = max(max_row, position[0]), max(max_col, position[1])
                        element.clear()
                    elif tag == SHEET_NS + "row":
                        element.clear()
            charts = 0
            for kind, drawing in relationships(part).values():
                if kind.endswith("/drawing") and drawing in names:
                    charts += sum(1 for chart_kind, _ in relationships(drawing).values() if chart_kind.endswith("/chart"))
            sheets.append({"name": name, "dimensions": dimension, "formulas": len(formulas), "first_formulas": formulas[:10], "grid": grid,
                           "size": (max_row, max_col), "charts": charts, **({} if complete else {"cellsChecked": MAX_CELLS, "complete": False})})
            if not complete:
                break
        # Shared strings only up to the highest index the first cells use.
        if wanted and strings_part in names:
            last, index = max(wanted), 0
            with package.open(strings_part) as stream:
                for _, element in iterparse(stream):
                    if element.tag == SHEET_NS + "si":
                        if index in wanted:
                            wanted[index] = "".join(node.text or "" for node in element.iter(SHEET_NS + "t"))
                        index += 1
                        element.clear()
                        if index > last:
                            break
    out = []
    for sheet in sheets:
        grid, (rows, columns) = sheet.pop("grid"), sheet.pop("size")
        shown = lambda value: clip(str(wanted.get(value[1]) or "")) if isinstance(value, tuple) else clip(str(value))
        sheet["first_rows"] = [[shown(grid[(row, column)]) if (row, column) in grid else None for column in range(1, columns + 1)] for row in range(1, rows + 1)]
        out.append(sheet)
    return {"kind": "xlsx", "sheets": out, "note": "formulas are shown as written; results are calculated only by Excel"}


def report(value):
    print(json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str))
