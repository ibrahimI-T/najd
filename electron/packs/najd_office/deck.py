"""Professional, editable decks with python-pptx: themes, layouts, native diagrams, charts, tables, images, Arabic and
right-to-left text, and PowerPoint animations/transitions (motion.py). Every element stays an editable slide object."""
import math

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import motion
from .text import is_arabic, set_rtl, set_run_font

THEMES = {
    # name: background, surface, text, muted, accents (3), heading font, body font
    "midnight": dict(bg="0B1F3A", surface="13294B", line="24426E", text="F2F5FA", muted="A9B6C8", accents=["22C3E6", "F5C542", "4CD97B"], heading="Segoe UI Semibold", body="Segoe UI"),
    "slate": dict(bg="FFFFFF", surface="F1F4F8", line="D5DCE5", text="1B2430", muted="5B6B7D", accents=["2F6FEB", "0FA37F", "E5534B"], heading="Segoe UI Semibold", body="Segoe UI"),
    "sand": dict(bg="F7F1E6", surface="FFFFFF", line="E2D6C2", text="2B2620", muted="6E655A", accents=["B5762F", "2F6F73", "8A4F7D"], heading="Segoe UI Semibold", body="Segoe UI"),
    "forest": dict(bg="0F2A22", surface="173A2F", line="2C5848", text="EEF5F0", muted="A7C1B3", accents=["7BD389", "F2C14E", "5BC0EB"], heading="Segoe UI Semibold", body="Segoe UI"),
}
SIZES = {"16:9": (13.333, 7.5), "4:3": (10.0, 7.5)}
CHARTS = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
          "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "stacked": XL_CHART_TYPE.COLUMN_STACKED}


class SlideRef:
    """A slide and its named shapes, so they can be animated by name."""

    def __init__(self, deck, slide):
        self.deck, self.slide, self.shapes = deck, slide, {}

    def animate(self, name, effect="fade", trigger="click", delay_ms=0, duration_ms=500, direction=None):
        self.deck.animate(self, name, effect, trigger, delay_ms, duration_ms, direction)
        return self

    def transition(self, kind="fade", duration_ms=700):
        self.deck.transition(self, kind, duration_ms)
        return self

    def notes(self, text):
        self.slide.notes_slide.notes_text_frame.text = str(text)
        return self


class Deck:
    def __init__(self, theme="midnight", size="16:9", direction=None, heading_font=None, body_font=None):
        if theme not in THEMES:
            raise ValueError("Unknown theme %r; use one of %s" % (theme, ", ".join(THEMES)))
        self.t = dict(THEMES[theme])
        if heading_font:
            self.t["heading"] = heading_font
        if body_font:
            self.t["body"] = body_font
        self.prs = Presentation()
        width, height = SIZES[size]
        self.prs.slide_width, self.prs.slide_height = Inches(width), Inches(height)
        self.w, self.h = width, height
        self.direction = direction  # "rtl", "ltr" or None (decided per text)
        self.margin = 0.6
        self._motion = {}
        self._transitions = {}

    # ---- basic building blocks ----
    def _rtl(self, text):
        return self.direction == "rtl" or (self.direction is None and is_arabic(text))

    def new_slide(self, background=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(background or self.t["bg"])
        return SlideRef(self, slide)

    def text(self, ref, text, x, y, w, h, size=20, bold=False, color=None, align=None, font=None, name=None, anchor="top", line_spacing=1.1):
        box = ref.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE  # fixed boxes: lint and the layout decide the fit
        frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for side in ("margin_left", "margin_right"):
            setattr(frame, side, Inches(0.08))
        frame.margin_top = frame.margin_bottom = Inches(0.04)
        lines = text if isinstance(text, (list, tuple)) else [text]
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.line_spacing = line_spacing
            run = paragraph.add_run()
            run.text = str(line)
            rtl = self._rtl(line)
            set_run_font(run, font or self.t["body"], size, bold, color or self.t["text"], arabic=is_arabic(line))
            if rtl:
                set_rtl(paragraph)
            if align:
                paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        if name:
            box.name = name
            ref.shapes[name] = box
        return box

    def box(self, ref, x, y, w, h, fill=None, line=None, radius=True, name=None, text=None, size=16, color=None, bold=False, shape=None):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        item = ref.slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            item.adjustments[0] = 0.12
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(fill or self.t["surface"])
        if line:
            item.line.color.rgb = RGBColor.from_string(line)
            item.line.width = Pt(1)
        else:
            item.line.fill.background()
        item.shadow.inherit = False
        if text is not None:
            frame = item.text_frame
            frame.word_wrap = True
            frame.margin_left = frame.margin_right = Inches(0.12)
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            lines = text if isinstance(text, (list, tuple)) else [text]
            for index, line_text in enumerate(lines):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.add_run()
                run.text = str(line_text)
                set_run_font(run, self.t["body"], size if index == 0 else max(10, size - 4), bold if index == 0 else False,
                             color or self.t["text"], arabic=is_arabic(line_text))
                if self._rtl(line_text):
                    set_rtl(paragraph)
                    paragraph.alignment = PP_ALIGN.CENTER
        if name:
            item.name = name
            ref.shapes[name] = item
        return item

    def connector(self, ref, start, end, color=None, name=None, arrow=True, begin_site=None, end_site=None):
        """A native connector glued to both shapes (moving a shape keeps the arrow attached)."""
        a, b = ref.shapes[start] if isinstance(start, str) else start, ref.shapes[end] if isinstance(end, str) else end
        line = ref.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
        sites = _sites(a, b)
        line.begin_connect(a, begin_site if begin_site is not None else sites[0])
        line.end_connect(b, end_site if end_site is not None else sites[1])
        line.line.color.rgb = RGBColor.from_string(color or self.t["muted"])
        line.line.width = Pt(2.5)
        if arrow:
            ln = line.line._get_or_add_ln()
            tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle", "w": "lg", "len": "lg"})
            ln.append(tail)
        if name:
            line.name = name
            ref.shapes[name] = line
        return line

    def area(self, subtitle=False):
        """Top and bottom (inches) of the content area under a heading."""
        return (2.05 if subtitle else 1.75), self.h - 0.55

    def centred(self, height, subtitle=False):
        """Top that centres a block of this height in the content area."""
        top, bottom = self.area(subtitle)
        return top + max(0.0, (bottom - top - height) / 2)

    def heading(self, ref, title, subtitle=None):
        rtl = self._rtl(title)
        self.text(ref, title, self.margin, 0.45, self.w - 2 * self.margin, 0.85, size=30, bold=True, font=self.t["heading"], name="title",
                  align="right" if rtl else "left")
        accent_x = self.w - self.margin - 1.2 if rtl else self.margin + 0.08
        bar = self.box(ref, accent_x, 1.3, 1.1, 0.06, fill=self.t["accents"][0], radius=False, name="accent")
        if subtitle:
            self.text(ref, subtitle, self.margin, 1.45, self.w - 2 * self.margin, 0.5, size=16, color=self.t["muted"], name="subtitle",
                      align="right" if rtl else "left")
        return bar

    # ---- layouts ----
    def title_slide(self, title, subtitle=None, kicker=None, image=None):
        ref = self.new_slide()
        rtl = self._rtl(title)
        if image:
            self.picture(ref, image, 0, 0, self.w, self.h, name="background")
            veil = self.box(ref, 0, 0, self.w, self.h, fill=self.t["bg"], radius=False, name="veil")
            _transparency(veil, 0.5)
        x, width = self.margin + 0.2, self.w - 2 * self.margin - 0.4
        if kicker:
            self.text(ref, kicker, x, 2.0, width, 0.5, size=16, bold=True, color=self.t["accents"][1], name="kicker", align="right" if rtl else "left")
        self.text(ref, title, x, 2.5, width, 1.6, size=44, bold=True, font=self.t["heading"], name="title", align="right" if rtl else "left", anchor="bottom")
        self.box(ref, self.w - self.margin - 1.6 if rtl else x + 0.08, 4.25, 1.4, 0.08, fill=self.t["accents"][0], radius=False, name="accent")
        if subtitle:
            self.text(ref, subtitle, x, 4.5, width, 1.0, size=20, color=self.t["muted"], name="subtitle", align="right" if rtl else "left")
        return ref

    def section(self, title, number=None, subtitle=None):
        ref = self.new_slide(self.t["surface"])
        rtl = self._rtl(title)
        if number is not None:
            self.text(ref, "%02d" % number if isinstance(number, int) else str(number), self.margin, 1.8, self.w - 2 * self.margin, 1.4, size=72, bold=True,
                      color=self.t["accents"][0], font=self.t["heading"], name="number", align="right" if rtl else "left")
        self.text(ref, title, self.margin, 3.3, self.w - 2 * self.margin, 1.2, size=38, bold=True, font=self.t["heading"], name="title", align="right" if rtl else "left")
        if subtitle:
            self.text(ref, subtitle, self.margin, 4.5, self.w - 2 * self.margin, 0.8, size=18, color=self.t["muted"], name="subtitle", align="right" if rtl else "left")
        return ref

    def bullets(self, title, items, subtitle=None, image=None):
        ref = self.new_slide()
        self.heading(ref, title, subtitle)
        rtl = self._rtl(title)
        top = 2.05 if subtitle else 1.75
        text_w = (self.w - 2 * self.margin) * (0.55 if image else 1.0)
        text_x = self.w - self.margin - text_w if (rtl and image) else self.margin
        for index, item in enumerate(items):
            y = top + index * 0.92
            dot_x = text_x + text_w - 0.22 if self._rtl(item) else text_x
            self.box(ref, dot_x, y + 0.2, 0.18, 0.18, fill=self.t["accents"][index % 3], radius=False, shape=MSO_SHAPE.OVAL, name="dot%d" % (index + 1))
            self.text(ref, item, text_x + (0 if self._rtl(item) else 0.35), y, text_w - 0.4, 0.8, size=20, name="point%d" % (index + 1),
                      align="right" if self._rtl(item) else "left")
        if image:
            image_x = self.margin if rtl else self.w - self.margin - (self.w - 2 * self.margin) * 0.4
            self.picture(ref, image, image_x, top, (self.w - 2 * self.margin) * 0.4, self.h - top - 0.7, name="image")
        return ref

    def columns(self, title, columns, subtitle=None):
        """Comparison cards: columns is a list of {"heading": str, "points": [str]} (2-4 columns). Cards fit their
        content and sit centred; each point has its own marker."""
        ref = self.new_slide()
        self.heading(ref, title, subtitle)
        count = len(columns)
        gap = 0.35
        width = (self.w - 2 * self.margin - gap * (count - 1)) / count
        most = max(len(column.get("points", [])) for column in columns)
        height = min(self.area(subtitle)[1] - self.area(subtitle)[0], 1.35 + most * 0.62 + 0.35)
        top = self.centred(height, subtitle)
        order = list(range(count))[::-1] if self._rtl(title) else list(range(count))
        for slot, index in enumerate(order):
            column = columns[index]
            rtl = self._rtl(column["heading"])
            x = self.margin + slot * (width + gap)
            accent = self.t["accents"][index % 3]
            self.box(ref, x, top, width, height, name="card%d" % (index + 1))
            self.box(ref, x + width - 0.35 - 0.9 if rtl else x + 0.35, top + 0.32, 0.9, 0.06, fill=accent, radius=False, name="cardbar%d" % (index + 1))
            self.text(ref, column["heading"], x + 0.3, top + 0.45, width - 0.6, 0.7, size=24, bold=True, font=self.t["heading"], name="cardtitle%d" % (index + 1),
                      align="right" if rtl else "left")
            lines = ["\u2022 " + point for point in column.get("points", [])]
            self.text(ref, lines, x + 0.3, top + 1.3, width - 0.6, height - 1.5, size=18, color=self.t["text"], name="cardtext%d" % (index + 1),
                      line_spacing=1.45, align="right" if rtl else "left")
        return ref

    def kpis(self, title, items, subtitle=None):
        """Big numbers: items is a list of (value, label); the cards sit centred in the content area."""
        ref = self.new_slide()
        self.heading(ref, title, subtitle)
        count = len(items)
        gap = 0.35
        width = (self.w - 2 * self.margin - gap * (count - 1)) / count
        height = 3.1
        top = self.centred(height, subtitle)
        order = list(range(count))[::-1] if self._rtl(title) else list(range(count))
        for slot, index in enumerate(order):
            value, label = items[index]
            x = self.margin + slot * (width + gap)
            accent = self.t["accents"][index % 3]
            self.box(ref, x, top, width, height, name="kpi%d" % (index + 1))
            self.text(ref, value, x, top + 0.45, width, 1.3, size=60, bold=True, color=accent, font=self.t["heading"], align="center", name="kpivalue%d" % (index + 1))
            self.box(ref, x + (width - 0.8) / 2, top + 1.85, 0.8, 0.05, fill=accent, radius=False, name="kpibar%d" % (index + 1))
            self.text(ref, label, x + 0.25, top + 2.05, width - 0.5, 0.85, size=18, color=self.t["text"], align="center", name="kpilabel%d" % (index + 1))
        return ref

    def timeline(self, title, steps, subtitle=None):
        """steps: list of (label, text). A native line with markers, centred in the content area; right to left for
        Arabic titles."""
        ref = self.new_slide()
        self.heading(ref, title, subtitle)
        rtl = self._rtl(title)
        count = len(steps)
        top, bottom = self.area(subtitle)
        y = (top + bottom) / 2
        self.box(ref, self.margin, y - 0.025, self.w - 2 * self.margin, 0.05, fill=self.t["line"], radius=False, name="axis")
        span = (self.w - 2 * self.margin) / count
        for index, (label, text) in enumerate(steps):
            slot = count - 1 - index if rtl else index
            cx = self.margin + span * (slot + 0.5)
            accent = self.t["accents"][index % 3]
            self.box(ref, cx - 0.2, y - 0.2, 0.4, 0.4, fill=accent, shape=MSO_SHAPE.OVAL, radius=False, name="marker%d" % (index + 1))
            self.text(ref, label, cx - span / 2 + 0.1, y - 1.35, span - 0.2, 0.9, size=22, bold=True, align="center", font=self.t["heading"], name="steplabel%d" % (index + 1), anchor="bottom")
            self.text(ref, text, cx - span / 2 + 0.15, y + 0.45, span - 0.3, 1.5, size=17, color=self.t["muted"], align="center", name="steptext%d" % (index + 1))
        return ref

    def diagram(self, title, nodes, edges=None, kind="flow", subtitle=None):
        """Native editable diagram. nodes: list of str or (title, detail); kind: flow, cycle or hierarchy (first node is
        the root). edges: list of (from_index, to_index); defaults to consecutive (and closing the cycle)."""
        ref = self.new_slide()
        self.heading(ref, title, subtitle)
        rtl = self._rtl(title)
        count = len(nodes)
        labels = [node if isinstance(node, (list, tuple)) else (node,) for node in nodes]
        top = 2.1 if subtitle else 1.9
        positions = []
        if kind == "flow":
            gap = 0.6
            width = min(2.8, (self.w - 2 * self.margin - gap * (count - 1)) / count)
            total = width * count + gap * (count - 1)
            start = (self.w - total) / 2
            flow_top = self.centred(1.8, subtitle)
            for index in range(count):
                slot = count - 1 - index if rtl else index
                positions.append((start + slot * (width + gap), flow_top, width, 1.8))
        elif kind == "cycle":
            radius_x, radius_y, cx, cy = 3.6, 1.95, self.w / 2, top + 2.45
            width, height = 2.3, 1.05
            for index in range(count):
                angle = -math.pi / 2 + (2 * math.pi * index / count) * (-1 if rtl else 1)
                positions.append((cx + radius_x * math.cos(angle) - width / 2, cy + radius_y * math.sin(angle) - height / 2, width, height))
        elif kind == "hierarchy":
            width, height = 2.6, 1.05
            positions.append(((self.w - width) / 2, top + 0.2, width, height))
            children = count - 1
            gap = 0.4
            child_w = min(2.6, (self.w - 2 * self.margin - gap * (children - 1)) / max(children, 1))
            total = child_w * children + gap * (children - 1)
            start = (self.w - total) / 2
            for index in range(children):
                slot = children - 1 - index if rtl else index
                positions.append((start + slot * (child_w + gap), top + 2.6, child_w, 1.3))
        else:
            raise ValueError("diagram kind must be flow, cycle or hierarchy")
        for index, (x, y, w, h) in enumerate(positions):
            accent = self.t["accents"][index % 3]
            self.box(ref, x, y, w, h, fill=self.t["surface"], line=accent, name="node%d" % (index + 1), text=list(labels[index]), size=17, bold=True)
        if edges is None:
            edges = [(0, index) for index in range(1, count)] if kind == "hierarchy" else [(index, index + 1) for index in range(count - 1)]
            if kind == "cycle" and count > 2:
                edges.append((count - 1, 0))
        for number, (a, b) in enumerate(edges, 1):
            self.connector(ref, "node%d" % (a + 1), "node%d" % (b + 1), name="link%d" % number)
        return ref

    def chart(self, title, categories, series, kind="column", subtitle=None, number_format=None, legend=True):
        """A native, editable chart. series: dict name -> values (same length as categories)."""
        if kind not in CHARTS:
            raise ValueError("chart kind must be one of %s" % ", ".join(CHARTS))
        ref = self.new_slide()
        self.heading(ref, title, subtitle)
        data = CategoryChartData()
        data.categories = list(categories)
        for name, values in series.items():
            data.add_series(name, list(values))
        top = 2.0 if subtitle else 1.8
        frame = ref.slide.shapes.add_chart(CHARTS[kind], Inches(self.margin), Inches(top), Inches(self.w - 2 * self.margin), Inches(self.h - top - 0.5), data)
        frame.name = "chart"
        ref.shapes["chart"] = frame
        chart = frame.chart
        chart.font.size = Pt(14)
        chart.font.color.rgb = RGBColor.from_string(self.t["muted"])
        chart.font.name = self.t["body"]
        chart.has_legend = legend and (len(series) > 1 or kind in ("pie", "doughnut"))
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        if kind in ("pie", "doughnut"):
            for index, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(self.t["accents"][index % 3] if index < 3 else self.t["line"])
            plot.has_data_labels = True
            plot.data_labels.font.color.rgb = RGBColor.from_string(self.t["text"])
        else:
            for index, item in enumerate(plot.series):
                fill = item.format.fill if kind != "line" else None
                if fill is not None:
                    fill.solid()
                    fill.fore_color.rgb = RGBColor.from_string(self.t["accents"][index % 3])
                else:
                    item.format.line.color.rgb = RGBColor.from_string(self.t["accents"][index % 3])
                    item.format.line.width = Pt(3)
            value_axis = chart.value_axis
            value_axis.has_major_gridlines = True
            value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(self.t["line"])
            value_axis.format.line.fill.background()
            if number_format:
                value_axis.tick_labels.number_format = number_format
                value_axis.tick_labels.number_format_is_linked = False
            chart.category_axis.format.line.color.rgb = RGBColor.from_string(self.t["line"])
            if self._rtl(title):
                # Right-to-left: the first category on the right, as Arabic readers expect.
                chart.category_axis.reverse_order = True
        return ref

    def table(self, title, header, rows, subtitle=None, widths=None):
        """A native table: header in the accent colour, banded rows, thin theme-coloured rules, text centred
        vertically with comfortable margins; right to left for Arabic."""
        ref = self.new_slide()
        self.heading(ref, title, subtitle)
        rtl = self._rtl(title)
        row_h = 0.72
        height = min(row_h * (len(rows) + 1), self.area(subtitle)[1] - self.area(subtitle)[0])
        top = self.centred(height, subtitle)
        frame = ref.slide.shapes.add_table(len(rows) + 1, len(header), Inches(self.margin), Inches(top), Inches(self.w - 2 * self.margin), Inches(height))
        frame.name = "table"
        ref.shapes["table"] = frame
        table = frame.table
        tbl_pr = table._tbl.tblPr
        tbl_pr.set("bandRow", "0"); tbl_pr.set("firstRow", "0")
        for child in list(tbl_pr):
            tbl_pr.remove(child)  # no built-in table style: the colours below are the whole look
        if rtl:
            tbl_pr.set("rtl", "1")
        if widths:
            total = sum(widths)
            for index, value in enumerate(widths):
                table.columns[index].width = Emu(int((self.w - 2 * self.margin) * 914400 * value / total))
        for r in range(len(rows) + 1):
            table.rows[r].height = Inches(height / (len(rows) + 1))
        for r, values in enumerate([header] + [list(row) for row in rows]):
            for c, value in enumerate(values):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(self.t["accents"][0] if r == 0 else (self.t["surface"] if r % 2 else self.t["bg"]))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left = cell.margin_right = Inches(0.18)
                _cell_rules(cell, self.t["line"])
                frame_text = cell.text_frame
                frame_text.paragraphs[0].text = ""
                run = frame_text.paragraphs[0].add_run()
                run.text = str(value)
                set_run_font(run, self.t["heading"] if r == 0 else self.t["body"], 18 if r == 0 else 17, r == 0,
                             (self.t["bg"] if r == 0 else self.t["text"]), arabic=is_arabic(value))
                if self._rtl(value) or rtl:
                    set_rtl(frame_text.paragraphs[0])
        return ref

    def picture(self, ref, path, x, y, w, h, name=None):
        """An image cropped to fill the box without distortion."""
        from PIL import Image
        with Image.open(path) as image:
            ratio = image.width / image.height
        target = w / h
        item = ref.slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        if ratio > target:
            cut = (1 - target / ratio) / 2
            item.crop_left = item.crop_right = cut
        elif ratio < target:
            cut = (1 - ratio / target) / 2
            item.crop_top = item.crop_bottom = cut
        if name:
            item.name = name
            ref.shapes[name] = item
        return item

    def image_slide(self, title, path, caption=None):
        ref = self.new_slide()
        self.heading(ref, title)
        self.picture(ref, path, self.margin, 1.75, self.w - 2 * self.margin, self.h - 2.6 if caption else self.h - 2.2, name="image")
        if caption:
            self.text(ref, caption, self.margin, self.h - 0.8, self.w - 2 * self.margin, 0.5, size=14, color=self.t["muted"], name="caption",
                      align="right" if self._rtl(caption) else "left")
        return ref

    def quote(self, text, author=None):
        ref = self.new_slide(self.t["surface"])
        rtl = self._rtl(text)
        self.text(ref, "”" if rtl else "“", self.margin, 0.9, 2, 1.6, size=120, bold=True, color=self.t["accents"][0], name="mark",
                  align="right" if rtl else "left", font="Georgia")
        self.text(ref, text, self.margin + 0.4, 2.3, self.w - 2 * self.margin - 0.8, 2.6, size=30, font=self.t["heading"], name="quote", align="right" if rtl else "left")
        if author:
            self.text(ref, author, self.margin + 0.4, 5.1, self.w - 2 * self.margin - 0.8, 0.6, size=18, color=self.t["muted"], name="author", align="right" if rtl else "left")
        return ref

    def closing(self, title, subtitle=None):
        ref = self.new_slide()
        self.text(ref, title, self.margin, 2.6, self.w - 2 * self.margin, 1.4, size=44, bold=True, font=self.t["heading"], align="center", name="title")
        self.box(ref, (self.w - 1.4) / 2, 4.1, 1.4, 0.08, fill=self.t["accents"][0], radius=False, name="accent")
        if subtitle:
            self.text(ref, subtitle, self.margin, 4.35, self.w - 2 * self.margin, 0.9, size=20, color=self.t["muted"], align="center", name="subtitle")
        return ref

    # ---- motion ----
    def animate(self, ref, name, effect="fade", trigger="click", delay_ms=0, duration_ms=500, direction=None):
        """Entrance effects appear, fade, fly, wipe, zoom; emphasis grow. trigger: click, with or after (previous)."""
        motion.check_effect(effect, trigger, delay_ms, duration_ms, direction)
        shape = ref.shapes[name] if isinstance(name, str) else name
        self._motion.setdefault(id(ref.slide), (ref.slide, []))[1].append(
            {"spid": shape.shape_id, "effect": effect, "trigger": trigger, "delay_ms": int(delay_ms), "duration_ms": int(duration_ms), "direction": direction})

    def reveal(self, ref, names, effect="fade", trigger="click", stagger_ms=150, duration_ms=500, direction=None):
        """Staged reveal: the first on the trigger, the rest one after another."""
        for index, name in enumerate(names):
            self.animate(ref, name, effect, trigger if index == 0 else "after", 0 if index == 0 else stagger_ms, duration_ms, direction)

    def transition(self, ref, kind="fade", duration_ms=700):
        if kind not in motion.TRANSITIONS:
            raise motion.MotionError("Unsupported transition %r; use fade, push, wipe or cut" % (kind,))
        self._transitions[id(ref.slide)] = (ref.slide, (kind, duration_ms))

    def save(self, path):
        touched = set(self._motion) | set(self._transitions)
        for key in touched:
            slide = (self._motion.get(key) or self._transitions.get(key))[0]
            items = self._motion.get(key, (None, []))[1]
            motion.apply(slide._element, items, self._transitions.get(key, (None, None))[1])
        self.prs.save(path)
        return path


def _sites(a, b):
    """Connection sites facing each other (rectangles: 0 top, 1 left, 2 bottom, 3 right)."""
    ax, ay = a.left + a.width / 2, a.top + a.height / 2
    bx, by = b.left + b.width / 2, b.top + b.height / 2
    if abs(bx - ax) >= abs(by - ay):
        return (3, 1) if bx > ax else (1, 3)
    return (2, 0) if by > ay else (0, 2)


def _cell_rules(cell, colour):
    """Thin rules in the theme's line colour on all four sides of a table cell."""
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in ("lnL", "lnR", "lnT", "lnB"):
        for old in tc_pr.findall(a + side):
            tc_pr.remove(old)
        line = tc_pr.makeelement(a + side, {"w": "9525", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        fill = line.makeelement(a + "solidFill", {})
        fill.append(fill.makeelement(a + "srgbClr", {"val": colour}))
        line.append(fill)
        tc_pr.insert(["lnL", "lnR", "lnT", "lnB"].index(side), line)


def _transparency(shape, alpha):
    fill = shape.fill._xPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
    color = fill[0] if fill is not None and len(fill) else None
    if color is not None:
        element = color.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", {"val": str(int((1 - alpha) * 100000))})
        color.append(element)
