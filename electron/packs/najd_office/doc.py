"""Word documents with python-docx: headings, paragraphs and tables with right-to-left Arabic handled."""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .text import is_arabic


class Doc:
    def __init__(self, font="Segoe UI", size=11, direction=None):
        self.document = docx.Document()
        self.direction = direction
        style = self.document.styles["Normal"]
        style.font.name = font
        style.font.size = Pt(size)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attribute), font)

    def _direction(self, paragraph, text):
        rtl = self.direction == "rtl" or (self.direction is None and is_arabic(text))
        if rtl:
            p_pr = paragraph._p.get_or_add_pPr()
            p_pr.append(p_pr.makeelement(qn("w:bidi"), {}))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                r_pr = run._r.get_or_add_rPr()
                r_pr.append(r_pr.makeelement(qn("w:rtl"), {}))
                lang = r_pr.makeelement(qn("w:lang"), {qn("w:bidi"): "ar-SA"})
                r_pr.append(lang)
        return paragraph

    def heading(self, text, level=1):
        return self._direction(self.document.add_heading(str(text), level), text)

    def paragraph(self, text, bold=False, size=None, color=None):
        paragraph = self.document.add_paragraph()
        run = paragraph.add_run(str(text))
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        return self._direction(paragraph, text)

    def bullets(self, items):
        for item in items:
            self._direction(self.document.add_paragraph(str(item), style="List Bullet"), item)

    def table(self, header, rows, style="Light Grid Accent 1"):
        table = self.document.add_table(rows=1, cols=len(header))
        try:
            table.style = style
        except Exception:
            pass
        for cell, value in zip(table.rows[0].cells, header):
            cell.text = str(value)
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = str(value)
        if self.direction == "rtl" or is_arabic(" ".join(map(str, header))):
            tbl_pr = table._tbl.tblPr
            tbl_pr.append(tbl_pr.makeelement(qn("w:bidiVisual"), {}))
        return table

    def image(self, path, width_in=6.0):
        return self.document.add_picture(path, width=Inches(width_in))

    def page_break(self):
        self.document.add_page_break()

    def save(self, path):
        self.document.save(path)
        return path
