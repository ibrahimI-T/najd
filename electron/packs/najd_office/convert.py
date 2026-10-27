"""Ready file conversions (owner request 2026-10-25): one call does what a model would otherwise write a script for.

    from najd_office import convert, report
    report(convert("photo.png", "jpg", width=1200))     # pictures: png jpg webp bmp gif ico pdf, optional width/height
    report(convert("table.csv", "xlsx"))                # csv -> xlsx (right to left when the text is Arabic), xlsx -> csv
    report(convert("report.docx", "md"))                # docx/pptx -> md or txt (headings, lists, tables, slide titles)
    report(convert(["a.png", "b.jpg"], "pdf"))          # several pictures -> one pdf

Every call returns {"ok", "output", "from", "to", ...} or {"ok": False, "error", "detail"}; nothing is overwritten unless
overwrite=True, and the output sits next to the input unless `output` names another file.
"""
import csv
import os
import re

IMAGE_OUT = {"png": "PNG", "jpg": "JPEG", "jpeg": "JPEG", "webp": "WEBP", "bmp": "BMP", "gif": "GIF", "ico": "ICO", "pdf": "PDF"}
IMAGE_IN = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".ico"}
ARABIC = re.compile(r"[\u0600-\u06FF]")


def _fail(error, detail=""):
    return {"ok": False, "error": error, "detail": str(detail)[:300]}


def _same(a, b):
    # Windows paths are case-insensitive and may be aliases: compare normalized paths, then the files themselves (C34-04).
    if os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b)):
        return True
    try:
        return os.path.exists(a) and os.path.exists(b) and os.path.samefile(a, b)
    except OSError:
        return True


def _target(source, to, output, overwrite):
    target = output or os.path.splitext(source)[0] + "." + to
    # The name must say what the bytes are (C36-01): jpg and jpeg are the same format.
    suffix = os.path.splitext(target)[1].lower().lstrip(".")
    if {suffix, to} != {"jpg", "jpeg"} and suffix != to:
        return None, _fail("WRONG_EXTENSION", f"{target} must end in .{to}")
    if _same(target, source):
        return None, _fail("SAME_FILE", target)
    if os.path.exists(target) and not overwrite:
        return None, _fail("EXISTS", target + " exists; pass overwrite=True or another output name")
    return target, None


def _images(sources, to, target, width, height, quality):
    from PIL import Image
    frames = []
    for source in sources:
        image = Image.open(source)
        image.load()
        if width or height:
            # Each picture keeps its own ratio: the requested size is never changed by an earlier picture (C34-03).
            w, h = image.size
            new_width = int(width) if width else max(1, round(w * int(height) / h))
            new_height = int(height) if height else max(1, round(h * int(width) / w))
            image = image.resize((new_width, new_height), Image.LANCZOS)
        if IMAGE_OUT[to] in ("JPEG", "PDF", "BMP") and image.mode not in ("RGB", "L"):
            image = image.convert("RGB")
        frames.append(image)
    options = {"quality": int(quality)} if IMAGE_OUT[to] in ("JPEG", "WEBP") else {}
    if IMAGE_OUT[to] == "PDF" and len(frames) > 1:
        frames[0].save(target, "PDF", save_all=True, append_images=frames[1:])
    else:
        frames[0].save(target, IMAGE_OUT[to], **options)
    return {"pages" if to == "pdf" else "size": len(frames) if to == "pdf" else list(frames[0].size)}


def _csv_to_xlsx(source, target):
    from openpyxl import Workbook
    with open(source, newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(handle, dialect))
    book = Workbook()
    sheet = book.active
    arabic = False
    for row in rows:
        values = []
        for cell in row:
            text = cell.strip()
            arabic = arabic or bool(ARABIC.search(text))
            try:
                values.append(int(text) if re.fullmatch(r"-?\d+", text) else float(text) if re.fullmatch(r"-?\d*\.\d+", text) else cell)
            except ValueError:
                values.append(cell)
        sheet.append(values)
        # Text stays text: a cell such as =WEBSERVICE(...) from the data never becomes a live formula (C34-02).
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.sheet_view.rightToLeft = arabic
    book.save(target)
    return {"rows": len(rows), "rightToLeft": arabic}


def _inert(value):
    # Text a spreadsheet could read as a formula starts with an apostrophe in the CSV, so opening it never runs anything (C38-01).
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _xlsx_to_csv(source, target, sheet):
    from openpyxl import load_workbook
    values = load_workbook(source, data_only=True, read_only=True)
    formulas = load_workbook(source, data_only=False, read_only=True)
    page = values[sheet] if sheet else values.worksheets[0]
    page_formulas = formulas[page.title]
    count = unvalued = 0
    with open(target, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        # A formula Excel never calculated has no stored value: its formula text is written and counted, never an empty cell (C37-01).
        for row, raw in zip(page.iter_rows(values_only=True), page_formulas.iter_rows(values_only=True)):
            out = []
            for value, formula in zip(row, raw):
                if value is None and isinstance(formula, str) and formula.startswith("="):
                    out.append(_inert(formula))
                    unvalued += 1
                else:
                    out.append("" if value is None else _inert(value))
            writer.writerow(out)
            count += 1
    result = {"rows": count, "sheet": page.title}
    if unvalued:
        result["formulasWithoutValue"] = unvalued
        result["note"] = "these cells hold formulas Excel never calculated; their formula text was written as plain text (it does not run)"
    return result


def _docx_text(source, markdown):
    import docx
    document = docx.Document(source)
    lines = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        level = re.search(r"heading (\d)", style)
        if markdown and level:
            lines.append("#" * int(level.group(1)) + " " + text)
        elif markdown and ("list" in style):
            lines.append("- " + text)
        else:
            lines.append(text)
    for table in document.tables:
        rows = [[cell.text.strip().replace("|", "/") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        # A table is one block: its rows stay on consecutive lines.
        if markdown:
            lines.append("\n".join(["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])] + ["| " + " | ".join(row) + " |" for row in rows[1:]]))
        else:
            lines.append("\n".join("\t".join(row) for row in rows))
    return "\n\n".join(lines) if markdown else "\n".join(lines)


def _pptx_text(source, markdown):
    from pptx import Presentation
    deck = Presentation(source)
    lines = []
    for number, slide in enumerate(deck.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        lines.append(("## " if markdown else "") + f"{number}. {title}".strip())
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(("- " if markdown else "  ") + text)
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def convert(source, to, output=None, overwrite=False, width=None, height=None, quality=90, sheet=None):
    """Convert one file (or several pictures into one pdf); see the module text for what converts to what."""
    to = str(to).lower().lstrip(".")
    sources = list(source) if isinstance(source, (list, tuple)) else [source]
    if not sources:
        return _fail("NO_INPUT")
    for item in sources:
        if not os.path.isfile(item):
            return _fail("NOT_FOUND", item)
    first = sources[0]
    kind = os.path.splitext(first)[1].lower()
    target, problem = _target(first, to, output, overwrite)
    if problem:
        return problem
    # Never onto any input, not only the first (C35-01).
    if any(_same(target, item) for item in sources):
        return _fail("SAME_FILE", target)
    try:
        if all(os.path.splitext(item)[1].lower() in IMAGE_IN for item in sources) and to in IMAGE_OUT:
            if len(sources) > 1 and to != "pdf":
                return _fail("ONE_INPUT", "several pictures convert only to one pdf")
            extra = _images(sources, to, target, width, height, quality)
        elif len(sources) > 1:
            return _fail("ONE_INPUT", "only pictures can be combined")
        elif kind == ".csv" and to == "xlsx":
            extra = _csv_to_xlsx(first, target)
        elif kind == ".xlsx" and to == "csv":
            extra = _xlsx_to_csv(first, target, sheet)
        elif kind in (".docx", ".pptx") and to in ("md", "txt"):
            text = (_docx_text if kind == ".docx" else _pptx_text)(first, to == "md")
            with open(target, "w", encoding="utf-8") as handle:
                handle.write(text)
            extra = {"characters": len(text)}
        elif kind in (".docx", ".xlsx", ".pptx") and to == "pdf":
            # Opening an Office file in another program needs the scanned, isolated review route; not done here (C34-01).
            return _fail("NOT_HERE", "Word, Excel or PowerPoint to pdf: call najd_office_review with pdf true (Office makes it from the checked copy)")
        else:
            return _fail("UNSUPPORTED", f"{kind or 'file'} -> {to}")
    except Exception as error:  # a damaged or unreadable file is a result, not a crash
        return _fail("CONVERT_FAILED", f"{type(error).__name__}: {error}")
    return {"ok": True, "output": os.path.abspath(target), "from": kind.lstrip("."), "to": to,
            "bytes": os.path.getsize(target), **extra}
