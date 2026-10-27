"""NAJD Office kit checks (docs/AGENT-PLATFORM-M3-PLAN.md acceptance for M3-7).
Run by tests/office-smoke.cjs with the engine's Python, the installed pack and the kit on sys.path:
    python -I office-kit-check.py <pack site-packages> <kit parent folder> <work folder>
Prints one line per check and KIT-OK at the end; any failure raises."""
import os
import site
import sys

site.addsitedir(sys.argv[1])
sys.path.append(sys.argv[2])
work = sys.argv[3]
os.chdir(work)

from najd_office import Deck, Doc, inspect, lint  # noqa: E402
from najd_office import motion  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def check(name, condition, detail=""):
    if not condition:
        raise AssertionError(name + (": " + str(detail) if detail else ""))
    print("ok", name)


# 1. Arabic paragraphs are right-to-left with complex-script fonts and the Arabic language tag.
deck = Deck(theme="midnight")
s = deck.title_slide("خطة الطاقة الشمسية للمزارع", "Solar pumping for farms", kicker="2026")
deck.save("arabic.pptx")
title = [shape for shape in Presentation("arabic.pptx").slides[0].shapes if shape.name == "title"][0]
paragraph = title.text_frame.paragraphs[0]
r_pr = paragraph.runs[0]._r.find(A + "rPr")
check("arabic rtl", paragraph._p.pPr.get("rtl") == "1")
check("arabic cs font", r_pr.find(A + "cs") is not None and r_pr.find(A + "cs").get("typeface") == "Segoe UI Semibold")
check("arabic lang", r_pr.get("lang") == "ar-SA")
check("arabic rPr order latin, ea, cs", [child.tag.split("}")[1] for child in r_pr if child.tag.split("}")[1] in ("latin", "ea", "cs")] == ["latin", "ea", "cs"])

# 2. Every offered effect and transition round-trips; unsupported ones are refused.
deck = Deck(theme="slate")
ref = deck.diagram("Flow", ["One", "Two", "Three", "Four", "Five", "Six"], kind="flow")
plan = [("node1", "appear", "click", None), ("node2", "fade", "with", None), ("node3", "fly", "after", "bottom"), ("node4", "wipe", "click", "left"),
        ("node5", "zoom", "after", None), ("node6", "grow", "click", None)]
for name, effect, trigger, direction in plan:
    ref.animate(name, effect, trigger, 100 if trigger == "after" else 0, 400, direction)
ref.transition("push", 600)
for kind in ("fade", "wipe", "cut"):
    deck.section("Section " + kind).transition(kind, 500)
deck.save("motion.pptx")
read = inspect("motion.pptx")
effects = read["slides"][0]["animations"]
check("effects round-trip", [(e["shape"], e["effect"], e["trigger"]) for e in effects] == [(n, e, t) for n, e, t, _ in plan], effects)
check("durations round-trip", all(e["duration_ms"] == 400 for e in effects if e["effect"] != "appear"), effects)
check("transitions round-trip", [slide["transition"]["kind"] for slide in read["slides"]] == ["push", "fade", "wipe", "cut"], [s["transition"] for s in read["slides"]])
for bad in (lambda: ref.animate("node1", "spin"), lambda: ref.animate("node1", "fade", "hover"), lambda: ref.animate("node1", "fade", direction="left"),
            lambda: ref.transition("morph")):
    try:
        bad()
    except motion.MotionError:
        continue
    raise AssertionError("unsupported motion accepted")
print("ok unsupported motion refused")

# 3. lint finds seeded issues and none on a clean deck.
clean = Deck(theme="midnight")
clean.title_slide("Clean title", "A subtitle that fits")
clean.bullets("Points", ["First point", "Second point", "Third point"])
clean.save("clean.pptx")
report = lint("clean.pptx")
check("clean deck has no issues", report["issues"] == [], report["issues"])
bad = Presentation()
bad.slide_width, bad.slide_height = Inches(13.333), Inches(7.5)
slide = bad.slides.add_slide(bad.slide_layouts[6])
tight = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.5)); tight.name = "tight"
tight.text_frame.word_wrap = True
tight.text_frame.text = "This sentence is far too long to fit inside such a small box on one line"
tight.text_frame.paragraphs[0].runs[0].font.size = Pt(20); tight.text_frame.paragraphs[0].runs[0].font.name = "Segoe UI"
over = slide.shapes.add_textbox(Inches(1.5), Inches(1.1), Inches(3), Inches(1)); over.name = "over"; over.text_frame.text = "Overlapping text"
off = slide.shapes.add_textbox(Inches(12.5), Inches(3), Inches(2), Inches(1)); off.name = "off"; off.text_frame.text = "Off slide"
faint = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(5), Inches(1)); faint.name = "faint"; faint.text_frame.text = "Faint text"
from pptx.dml.color import RGBColor  # noqa: E402
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
faint.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("DDDDDD")
font = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(5), Inches(1)); font.name = "font"; font.text_frame.text = "Missing font"
font.text_frame.paragraphs[0].runs[0].font.name = "NoSuchFontNajd"
arabic = slide.shapes.add_textbox(Inches(7), Inches(5.5), Inches(5), Inches(1)); arabic.name = "arabicltr"; arabic.text_frame.text = "نص عربي"
arabic.text_frame.paragraphs[0]._p.get_or_add_pPr()
bad.save("bad.pptx")
# A dangling animation target.
from lxml import etree  # noqa: E402
motion_bad = Presentation("bad.pptx")
motion.apply(motion_bad.slides[0]._element, [{"spid": 999, "effect": "fade", "trigger": "click", "delay_ms": 0, "duration_ms": 300, "direction": None}])
motion_bad.save("bad.pptx")
found = [issue["issue"] for issue in lint("bad.pptx")["issues"]]
for expected in ("text overflows", "text overlaps", "off-slide", "low contrast", "font not installed: NoSuchFontNajd", "Arabic paragraph not right-to-left", "animation targets a missing shape"):
    check("lint finds " + expected, any(item.startswith(expected) for item in found), found)

# 4. inspect reads documents and workbooks too.
doc = Doc()
doc.heading("التقرير السنوي")
doc.paragraph("Revenue grew in every region.")
doc.table(["المنطقة", "المبيعات"], [["الرياض", 120]])
doc.save("report.docx")
docx_view = inspect("report.docx")
check("docx inspect rtl", docx_view["paragraphs"][0].get("rtl") is True and docx_view["tables"] == [[2, 2]], docx_view)
import openpyxl  # noqa: E402
book = openpyxl.Workbook(); sheet = book.active; sheet["A1"] = 2; sheet["A2"] = 3; sheet["A3"] = "=SUM(A1:A2)"; book.save("sum.xlsx")
xlsx_view = inspect("sum.xlsx")
check("xlsx inspect formulas", xlsx_view["sheets"][0]["formulas"] == 1 and "calculated only by Excel" in xlsx_view["note"])
# M3-I019: a sparse workbook (A1 and XFD1048576) is inspected by its cells, not its declared rectangle.
import time  # noqa: E402
book = openpyxl.Workbook(); sheet = book.active; sheet["A1"] = "=1+1"; sheet["XFD1048576"] = 5; book.save("sparse.xlsx")
started = time.time(); sparse_view = inspect("sparse.xlsx")
check("xlsx sparse inspect bounded", time.time() - started < 10 and sparse_view["sheets"][0]["formulas"] == 1 and len(sparse_view["sheets"][0]["first_rows"]) <= 8 and all(len(row) <= 26 for row in sparse_view["sheets"][0]["first_rows"]), sparse_view)
# M3-I023: a whole-sheet merged range and a whole-sheet hyperlink range do not expand while the workbook is read.
# Written as raw XML: building them with openpyxl would expand them while writing.
import zipfile  # noqa: E402
with zipfile.ZipFile("ranges.xlsx", "w") as package:
    package.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
    package.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    package.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Big" sheetId="1" r:id="s1"/></sheets></workbook>')
    package.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="s1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
    package.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><dimension ref="A1:XFD1048576"/><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>top</t></is></c><c r="B1"><f>SUM(C1:C2)</f><v>0</v></c></row></sheetData>'
                     '<mergeCells count="1"><mergeCell ref="A2:XFD1048576"/></mergeCells><hyperlinks><hyperlink ref="A1:XFD1048576" location="Big!A1"/></hyperlinks></worksheet>')
started = time.time(); ranges_view = inspect("ranges.xlsx")
check("xlsx merged and hyperlink ranges bounded", time.time() - started < 10 and ranges_view["sheets"][0]["formulas"] == 1 and ranges_view["sheets"][0]["first_rows"] == [["top", "=SUM(C1:C2)"]], ranges_view)
# M3-I026: a paragraph that is one long word wraps by characters; text in a box with wrapping off is checked for width.
from najd_office.text import measure_lines  # noqa: E402
url = "https://example.com/" + "a" * 120
check("long first word wraps", measure_lines(url, "Segoe UI", 18, 120) >= 5, measure_lines(url, "Segoe UI", 18, 120))
narrow = Presentation(); narrow_slide = narrow.slides.add_slide(narrow.slide_layouts[6])
for wrap, label in ((True, "wrapped"), (False, "unwrapped")):
    box = narrow_slide.shapes.add_textbox(Inches(0.5 if wrap else 5), Inches(0.5), Inches(1.5), Inches(0.6)); box.name = label
    box.text_frame.word_wrap = wrap; box.text_frame.text = url; box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
narrow.save("narrow.pptx")
narrow_issues = {(issue["shape"], issue["issue"].split(" (")[0]) for issue in lint("narrow.pptx")["issues"]}
check("long URL overflow found", ("wrapped", "text overflows") in narrow_issues and ("unwrapped", "text wider than its box") in narrow_issues, narrow_issues)

# Ready conversions (owner request 2026-10-25): one call each, nothing overwritten, a table stays one Markdown block.
from najd_office import convert
from PIL import Image as _Image
_Image.new("RGBA", (800, 400), (200, 50, 50, 128)).save("conv.png")
_Image.new("RGB", (300, 300), (20, 120, 200)).save("conv2.jpg")
picture = convert("conv.png", "jpg", width=400)
check("png to jpg resized", picture["ok"] and picture["size"] == [400, 200], picture)
check("an existing output is not overwritten", convert("conv.png", "jpg")["error"] == "EXISTS")
check("two pictures to one pdf", convert(["conv.png", "conv2.jpg"], "pdf", output="conv-both.pdf").get("pages") == 2)
with open("conv.csv", "w", encoding="utf-8") as handle:
    handle.write("البند,المبلغ\nإيجار,2500\n")
sheet = convert("conv.csv", "xlsx")
check("csv to xlsx, right to left for Arabic", sheet["ok"] and sheet["rightToLeft"] is True and sheet["rows"] == 2, sheet)
check("xlsx back to csv", convert("conv.xlsx", "csv", output="conv-back.csv")["rows"] == 2)
import docx as _docx
_document = _docx.Document(); _document.add_heading("عنوان", 1)
_table = _document.add_table(rows=2, cols=2); _table.cell(0, 0).text = "أ"; _table.cell(0, 1).text = "ب"; _table.cell(1, 0).text = "1"; _table.cell(1, 1).text = "2"
_document.save("conv.docx")
markdown = convert("conv.docx", "md")
text = open("conv.md", encoding="utf-8").read()
check("docx to markdown with one table block", markdown["ok"] and text.startswith("# عنوان") and "| أ | ب |\n|---|---|\n| 1 | 2 |" in text, text)
check("unsupported pair refused", convert("conv.csv", "pptx")["error"] == "UNSUPPORTED")
check("lint gives a note for a document", "note" in lint("conv.docx"))
with open("risky.csv", "w", encoding="utf-8") as handle:
    handle.write('name,link\nx,=WEBSERVICE("https://example.invalid/")\n')
convert("risky.csv", "xlsx")
from openpyxl import load_workbook as _load
_cell = _load("risky.xlsx")["Sheet"]["B2"]
check("a formula in csv data stays text", _cell.data_type == "s" and str(_cell.value).startswith("="), (_cell.data_type, _cell.value))
_Image.new("RGB", (200, 400)).save("tall.png"); _Image.new("RGB", (400, 200)).save("wide.png")
_mixed = convert(["tall.png", "wide.png"], "pdf", width=100, output="mixed.pdf")
check("a tall and a wide picture both convert into one pdf at the same width", _mixed["ok"] and _mixed["pages"] == 2, _mixed)
_tall = convert("tall.png", "jpg", width=100); _wide = convert("wide.png", "jpg", width=100)
check("resize by width keeps the ratio", _tall["size"] == [100, 200] and _wide["size"] == [100, 50], (_tall, _wide))
check("the source cannot be overwritten through a different letter case", convert("conv.csv", "csv", output="CONV.CSV", overwrite=True)["error"] == "SAME_FILE")
check("Office to pdf is left to the review route", convert("conv.docx", "pdf")["error"] == "NOT_HERE")
from openpyxl import Workbook as _Workbook
_book = _Workbook(); _book.active.append([1]); _book.active.append([2]); _book.active.append(["=SUM(A1:A2)"]); _book.save("formula.xlsx")
_csv = convert("formula.xlsx", "csv")
check("an uncalculated formula is written, not lost", _csv["formulasWithoutValue"] == 1 and "'=SUM(A1:A2)" in open("formula.csv", encoding="utf-8-sig").read(), _csv)
check("an output name must match its format", convert("tall.png", "jpg", output="result.png")["error"] == "WRONG_EXTENSION" and convert("tall.png", "jpeg", output="tall2.jpg")["ok"])
check("a batch never writes onto a later input", convert(["tall.png", "wide.png"], "pdf", output="wide.png", overwrite=True)["error"] in ("SAME_FILE", "WRONG_EXTENSION") and os.path.getsize("wide.png") > 0)
print("KIT-OK")
